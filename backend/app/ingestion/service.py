"""Ingestion service for parsing and registering business data files into DuckDB."""
import csv
import io
import re
import uuid
from pathlib import Path
from typing import Dict, List, Optional
from app.core.config import settings
from app.core.database import db_manager
from app.ingestion.models import ColumnMetadata, DatasetInfo


def sanitize_column_name(name: str) -> str:
    """Sanitize column name to be clean, lowercase, alphanumeric with underscores."""
    clean = re.sub(r"[^\w\s]", "_", str(name).strip())
    clean = re.sub(r"\s+", "_", clean)
    clean = clean.strip("_").lower()
    if not clean or clean[0].isdigit():
        clean = f"col_{clean}"
    return clean


class IngestionService:
    """Service handling file upload, parsing, and DuckDB registration."""

    def __init__(self):
        self.upload_dir = Path(settings.UPLOAD_DIR)
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        self._dataset_registry: Dict[str, DatasetInfo] = {}

    def get_all_datasets(self) -> List[DatasetInfo]:
        """Return all registered datasets."""
        return list(self._dataset_registry.values())

    def get_dataset(self, dataset_id: str) -> Optional[DatasetInfo]:
        """Get dataset by ID."""
        return self._dataset_registry.get(dataset_id)

    async def process_file(
        self,
        filename: str,
        content: bytes,
        sheet_name: Optional[str] = None,
        custom_name: Optional[str] = None,
    ) -> DatasetInfo:
        """Process and register uploaded dataset into DuckDB."""
        dataset_id = str(uuid.uuid4())[:8]
        safe_base = Path(filename).stem
        safe_base = re.sub(r"[^\w\-]", "_", safe_base)
        table_name = f"dataset_{dataset_id}_{safe_base.lower()}"

        file_ext = Path(filename).suffix.lower()
        saved_file_path = self.upload_dir / f"{dataset_id}_{filename}"
        saved_file_path.write_bytes(content)
        path_str = str(saved_file_path.resolve()).replace("\\", "/")

        if file_ext in [".csv", ".txt"]:
            # DuckDB direct read with auto detection
            db_manager.connection.execute(
                f'CREATE OR REPLACE TABLE "{table_name}" AS SELECT * FROM read_csv_auto(\'{path_str}\', header=True)'
            )
        elif file_ext == ".parquet":
            db_manager.connection.execute(
                f'CREATE OR REPLACE TABLE "{table_name}" AS SELECT * FROM read_parquet(\'{path_str}\')'
            )
        elif file_ext in [".xlsx", ".xls"]:
            # Openpyxl parser
            import openpyxl
            wb = openpyxl.load_workbook(saved_file_path, data_only=True)
            ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                raise ValueError("Excel file is empty.")
            header = rows[0]
            data_rows = rows[1:]
            
            # Write to temporary CSV for DuckDB loading
            csv_path = self.upload_dir / f"{dataset_id}_converted.csv"
            with open(csv_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(header)
                writer.writerows(data_rows)
            csv_path_str = str(csv_path.resolve()).replace("\\", "/")
            db_manager.connection.execute(
                f'CREATE OR REPLACE TABLE "{table_name}" AS SELECT * FROM read_csv_auto(\'{csv_path_str}\', header=True)'
            )
        else:
            raise ValueError(f"Unsupported file format: {file_ext}")

        # Get table schema from DuckDB
        describe_rows = db_manager.get_table_schema(table_name)
        
        # Rename columns in DuckDB table to sanitized names if needed
        columns_meta: List[ColumnMetadata] = []
        seen = set()
        rename_statements = []

        for desc in describe_rows:
            raw_col = desc["column_name"]
            clean_col = sanitize_column_name(raw_col)
            suffix = 1
            unique_clean = clean_col
            while unique_clean in seen:
                unique_clean = f"{clean_col}_{suffix}"
                suffix += 1
            seen.add(unique_clean)

            if unique_clean != raw_col:
                rename_statements.append(f'ALTER TABLE "{table_name}" RENAME COLUMN "{raw_col}" TO "{unique_clean}"')

            # Fetch sample values
            samples_res = db_manager.execute_query(
                f'SELECT "{raw_col}" FROM "{table_name}" WHERE "{raw_col}" IS NOT NULL LIMIT 3'
            )
            samples = [list(r.values())[0] for r in samples_res]

            columns_meta.append(
                ColumnMetadata(
                    name=unique_clean,
                    original_name=raw_col,
                    data_type=str(desc["column_type"]),
                    sample_values=samples,
                )
            )

        for stmt in rename_statements:
            db_manager.connection.execute(stmt)

        # Get total row count and preview rows
        row_count = int(db_manager.execute_scalar(f'SELECT count(*) FROM "{table_name}"') or 0)
        preview_rows = db_manager.execute_query(f'SELECT * FROM "{table_name}" LIMIT 10')

        dataset_info = DatasetInfo(
            id=dataset_id,
            name=custom_name or filename,
            filename=filename,
            file_type=file_ext.lstrip("."),
            file_size_bytes=len(content),
            row_count=row_count,
            column_count=len(columns_meta),
            columns=columns_meta,
            table_name=table_name,
            preview_rows=preview_rows,
        )

        self._dataset_registry[dataset_id] = dataset_info
        return dataset_info

    def delete_dataset(self, dataset_id: str) -> bool:
        """Delete dataset from registry and drop DuckDB table."""
        if dataset_id in self._dataset_registry:
            info = self._dataset_registry.pop(dataset_id)
            try:
                db_manager.connection.execute(f'DROP TABLE IF EXISTS "{info.table_name}"')
            except Exception:
                pass
            return True
        return False


ingestion_service = IngestionService()
